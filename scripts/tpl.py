"""
Component library that reproduces the ACS-UPM template exactly.

Every measurement, colour, font and shape below is taken from template.pptx:
palette and typography from its slides, rounded-corner value adj=4368 from its
cards, 0.45 in circles for the numbered badges, footer at y=7.06 in 9 pt
Helvetica #6B7683, logo at x=9.92 y=0.21 w=3.10 h=0.69.
"""
import copy
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PARENT = os.path.dirname(HERE)


def _find(*rel):
    """Look for the template assets next to this file or one level up."""
    for base in (HERE, PARENT):
        p = os.path.join(base, *rel)
        if os.path.exists(p):
            return p
    return os.path.join(HERE, *rel)


TEMPLATE = _find("template.pptx")
MEDIA = _find("tpl_media", "ppt", "media")

# ---------------------------------------------------------------- palette
PURPLE = "964BFF"
TEAL = "42DEDC"
TEAL_D = "08A3A9"
NAVY = "23263B"
NAVY2 = "24263A"
GREY_TXT = "6B7683"
LILAC = "E1CAFF"
GREY1 = "F4F4F6"
GREY2 = "F3F3F7"
GREY3 = "F4F6FA"
CODE_TXT = "E8EDF2"
WHITE = "FFFFFF"

# ---------------------------------------------------------------- fonts
NOHEMI = "Nohemi"
HN = "Helvetica Neue"
HNL = "Helvetica Neue Light"
HNM = "Helvetica Neue Medium"
HELV = "Helvetica"
MONO = "Courier New"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

FOOTER_CENTRE = "Module 2. S01: Data models"
FOOTER_LEFT = "ACS-UPM Diploma"


class Deck:
    def __init__(self):
        self.prs = Presentation(TEMPLATE)
        src = self.prs.slides
        self._cover_bg = copy.deepcopy(src[0]._element.find(".//" + P + "bg"))
        self._content_bg = copy.deepcopy(src[1]._element.find(".//" + P + "bg"))
        # remove the eight example slides, keep master, layout, theme and media
        xml_slides = self.prs.slides._sldIdLst
        for sld in list(xml_slides):
            rId = sld.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            self.prs.part.drop_rel(rId)
            xml_slides.remove(sld)
        self.layout = self.prs.slide_masters[0].slide_layouts[0]
        self.n = 0

    # ------------------------------------------------------------ slides
    def _blank(self, cover=False):
        s = self.prs.slides.add_slide(self.layout)
        self._last = s
        csld = s._element.find(P + "cSld")
        old = csld.find(P + "bg")
        if old is not None:
            csld.remove(old)
        bg = copy.deepcopy(self._cover_bg if cover else self._content_bg)
        csld.insert(0, bg)
        if cover:
            # the copied background points at the template's own relationship id;
            # register the image in this slide and rewrite the reference
            blip = bg.find(".//" + A + "blip")
            _, rId = s.part.get_or_add_image_part(os.path.join(MEDIA, "image1.png"))
            blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rId)
        return s

    def cover(self, session, line1, line2, sub1, sub2, author, year, ribbon=None):
        s = self._blank(cover=True)
        veil = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0, 0,
                                  self.prs.slide_width, self.prs.slide_height)
        _round(veil, 0)
        _fill(veil, PURPLE, alpha=40)
        _noline(veil)
        s.shapes.add_picture(os.path.join(MEDIA, "image2.png"),
                             Inches(8.97), Inches(0.21), Inches(4.04), Inches(0.89))
        txt(s, session, 0.90, 1.20, 11.50, 0.50, HNL, 20, NAVY)
        txt(s, line1, 0.90, 1.80, 11.50, 0.90, NOHEMI, 44, NAVY, bold=True)
        txt(s, line2, 0.90, 2.70, 11.50, 0.90, NOHEMI, 44, NAVY, bold=True)
        b = s.shapes.add_textbox(Inches(0.90), Inches(3.85), Inches(11.50), Inches(0.80))
        tf = b.text_frame
        tf.word_wrap = True
        _run(tf.paragraphs[0], sub1, HNL, 15, NAVY)
        p2 = tf.add_paragraph()
        _run(p2, sub2, HN, 15, NAVY, bold=True)
        txt(s, author, 0.90, 6.55, 4.50, 0.40, "Arial", 14, NAVY)
        txt(s, year, 11.40, 6.55, 1.20, 0.40, "Arial", 14, NAVY, align=PP_ALIGN.RIGHT)
        if ribbon:
            r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.90), Inches(5.39),
                                   Inches(5.54), Inches(0.62))
            _round(r, 10000)
            _fill(r, WHITE, alpha=55)
            _line(r, WHITE, 1)
            txt(s, ribbon, 1.10, 5.43, 5.38, 0.55, HNM, 15, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        self.n += 1
        return s

    def slide(self, title, title_w=9.40):
        s = self._blank()
        s.shapes.add_picture(os.path.join(MEDIA, "image3.png"),
                             Inches(9.92), Inches(0.21), Inches(3.10), Inches(0.69))
        txt(s, title, 0.41, 0.27, title_w, 0.59, NOHEMI, 28, NAVY, bold=True)
        self.n += 1
        txt(s, FOOTER_LEFT, 0.41, 7.06, 2.60, 0.15, HELV, 9, GREY_TXT)
        txt(s, FOOTER_CENTRE, 3.32, 7.06, 6.60, 0.15, HELV, 9, GREY_TXT, align=PP_ALIGN.CENTER)
        txt(s, str(self.n), 12.43, 7.06, 0.50, 0.15, HELV, 9, GREY_TXT, align=PP_ALIGN.RIGHT)
        return s

    def section(self, kicker, title, lead=None, chips=None):
        """Dark divider, in the template's navy."""
        s = self._blank()
        band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                  self.prs.slide_width, self.prs.slide_height)
        _fill(band, NAVY)
        _noline(band)
        s.shapes.add_picture(os.path.join(MEDIA, "image2.png"),
                             Inches(9.92), Inches(0.21), Inches(3.10), Inches(0.69))
        txt(s, kicker, 0.90, 2.40, 6.0, 0.40, HN, 16, TEAL, bold=True)
        txt(s, title, 0.90, 2.85, 11.5, 1.10, NOHEMI, 40, WHITE, bold=True)
        if lead:
            txt(s, lead, 0.90, 4.15, 11.0, 0.90, HNL, 17, "C9CEDA")
        if chips:
            x = 0.90
            for c in chips:
                w = 0.115 * len(c) + 0.34
                sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.35),
                                        Inches(w), Inches(0.36))
                _round(sh, 20000)
                _fill(sh, "39405A")
                _noline(sh)
                txt(s, c, x, 5.35, w, 0.36, MONO, 11, WHITE,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                x += w + 0.14
        self.n += 1
        return s

    def notes(self, s, text):
        ns = s.notes_slide
        tf = ns.notes_text_frame
        if tf is None:                     # the template ships no notes placeholder
            from pptx.oxml.ns import qn
            sp_tree = ns._element.find(qn("p:cSld")).find(qn("p:spTree"))
            xml = (
                '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '<p:nvSpPr><p:cNvPr id="99" name="Notes Placeholder"/>'
                '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
                '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
                '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
            )
            sp_tree.append(etree.fromstring(xml))
            tf = ns.notes_text_frame
        tf.text = text

    def save(self, path):
        self.prs.save(path)
        return path


# ---------------------------------------------------------------- primitives
def _fill(shape, hexcolor, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hexcolor)
    if alpha is not None:
        srgb = shape.fill._xPr.find(A + "solidFill").find(A + "srgbClr")
        a = etree.SubElement(srgb, A + "alpha")
        a.set("val", str(int(alpha * 1000)))


def _noline(shape):
    shape.line.fill.background()


def _line(shape, hexcolor, pt):
    shape.line.color.rgb = RGBColor.from_string(hexcolor)
    shape.line.width = Pt(pt)


def _round(shape, adj):
    gd = shape._element.spPr.find(A + "prstGeom").find(A + "avLst").find(A + "gd")
    if gd is None:
        gd = etree.SubElement(shape._element.spPr.find(A + "prstGeom").find(A + "avLst"), A + "gd")
        gd.set("name", "adj")
    gd.set("fmla", f"val {adj}")


def _run(par, text, font, size, color, bold=False, italic=False):
    r = par.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def txt(s, text, x, y, w, h, font, size, color, bold=False, italic=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, wrap=True):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(str(text).split("\n")):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        if spacing:
            par.line_spacing = spacing
        _run(par, line, font, size, color, bold, italic)
    return box


def rich(s, parts, x, y, w, h, size=14.5, color=NAVY, spacing=1.15, space_after=6,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """parts: list of paragraphs; each paragraph is a list of (text, bold) tuples."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(parts):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.line_spacing = spacing
        par.space_after = Pt(space_after)
        if isinstance(para, str):
            para = [(para, False)]
        for text, bold in para:
            _run(par, text, HN if bold else HNL, size, color, bold)
    return box


def card(s, x, y, w, h, color=LILAC, alpha=40):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    _round(sh, 4368)
    _fill(sh, color, alpha=alpha)
    _noline(sh)
    return sh


def badge(s, number, x, y, color=PURPLE, d=0.45):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    _fill(sh, color)
    _noline(sh)
    txt(s, str(number), x, y, d, d, HN, 14, WHITE, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return sh


def point(s, number, heading, body, x, y, w=5.40, color=PURPLE, gap=0.54, size=14.5):
    """The template's numbered point: circle, bold heading, light body."""
    badge(s, number, x, y - 0.02, color)
    txt(s, heading, x + gap, y + 0.03, w - gap, 0.32, HN, 16, NAVY, bold=True)
    return rich(s, [body], x + gap, y + 0.37, w - gap, 1.0, size=size)


def code(s, text, x, y, w, h, size=12.5, fill=NAVY):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    _round(sh, 3500)
    _fill(sh, fill)
    _noline(sh)
    txt(s, text, x + 0.18, y + 0.14, w - 0.36, h - 0.28, MONO, size, CODE_TXT, spacing=1.12)
    return sh


def pill(s, text, x, y, w=None, color=TEAL_D, size=9, h=0.33, txtcolor=WHITE):
    w = w or (0.085 * len(text) + 0.34)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    _round(sh, 22000)
    _fill(sh, color)
    _noline(sh)
    txt(s, text, x, y, w, h, HN, size, txtcolor, bold=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def pillrow(s, items, x, y, gap=0.10, **kw):
    cx = x
    for it in items:
        cx += pill(s, it, cx, y, **kw) + gap
    return cx


def rule(s, x, y, w, color=TEAL_D, pt=1.5):
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(pt))
    _fill(ln, color)
    _noline(ln)
    return ln


def banner(s, label, text, x=2.15, y=5.75, w=9.05, h=0.52, fill=NAVY2,
           label_color=WHITE, text_color=WHITE, size=12.5):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    _round(sh, 12000)
    _fill(sh, fill)
    _noline(sh)
    box = s.shapes.add_textbox(Inches(x + 0.20), Inches(y), Inches(w - 0.40), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par = tf.paragraphs[0]
    par.line_spacing = 1.1
    if label:
        _run(par, label + " ", HN, size, label_color, bold=True)
    _run(par, text, HNL, size, text_color)
    return sh


def table(s, rows, x, y, w, colw, rowh=0.40, size=11, head_fill=NAVY,
          bold_col=None, mono_cols=(), mono_size=9.5, row_fills=None, hot_col=None):
    n_r, n_c = len(rows), len(rows[0])
    shp = s.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(rowh * n_r))
    tbl = shp.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, cw in enumerate(colw):
        tbl.columns[j].width = Inches(cw)
    for i in range(n_r):
        tbl.rows[i].height = Inches(rowh if i else rowh * 1.05)
        for j in range(n_c):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = RGBColor.from_string(head_fill)
            elif row_fills:
                cell.fill.fore_color.rgb = RGBColor.from_string(row_fills(i - 1, j))
            else:
                cell.fill.fore_color.rgb = RGBColor.from_string(GREY1 if i % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            is_mono = j in mono_cols
            bold = (i == 0) or (bold_col is not None and j == bold_col)
            colr = WHITE if i == 0 else (PURPLE if (hot_col is not None and j == hot_col) else NAVY)
            _run(par, str(rows[i][j]),
                 MONO if (is_mono and i > 0) else (HN if bold else HNL),
                 mono_size if (is_mono and i > 0) else size, colr, bold)
    return shp


def session_strip(s, total, highlight, y=1.38, x0=0.41, x1=12.93, h=0.34):
    """The template's strip of numbered sessions: highlighted ones in purple."""
    pitch = (x1 - x0) / total
    w = pitch - 0.06
    for i in range(1, total + 1):
        x = x0 + (i - 1) * pitch
        on = i in highlight
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        _round(sh, 14000)
        _fill(sh, PURPLE if on else GREY1)
        _noline(sh)
        txt(s, str(i), x, y, w, h, HN, 11, WHITE if on else "8A93A0", bold=True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def timeline(s, cols, axis_y=4.24, x0=0.53, pitch=2.16, colw=1.80, marker=None):
    """The template's horizontal course timeline.

    cols: list of dicts with keys kicker, title, date, meta, body and optional tag.
    marker: index (0-based) of the column to mark as the current session.
    """
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(axis_y), Inches(13.333), Pt(1))
    _fill(line, "3A3D52")
    _noline(line)
    for i, c in enumerate(cols):
        x = x0 + i * pitch
        tick = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.02), Inches(axis_y - 0.10),
                                  Pt(1), Inches(0.22))
        _fill(tick, "3A3D52")
        _noline(tick)
        if c.get("tag"):
            txt(s, c["tag"], x, 2.21, colw + 0.6, 0.26, HNL, 13, PURPLE)
        txt(s, c["kicker"], x, 2.62, colw, 0.26, HNL, 12, NAVY)
        txt(s, c["title"], x, 2.88, colw, 0.95, HN, 13, NAVY, bold=True, spacing=1.08)
        txt(s, c["date"], x, 4.58, colw + 0.25, 0.28, HNL, 14, PURPLE)
        txt(s, c["meta"], x, 4.86, colw + 0.25, 0.26, HNL, 9.5, GREY_TXT)
        txt(s, c["body"], x, 5.28, colw, 1.60, HNL, 10.5, GREY_TXT, spacing=1.12)
    if marker is not None:
        mx = x0 + marker * pitch
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(mx - 0.10), Inches(axis_y - 0.13),
                                 Inches(0.26), Inches(0.26))
        _fill(dot, PURPLE)
        _line(dot, WHITE, 1.75)


def timeline_sessions(s, cols, axis_y=2.80, x0=0.50, pitch=2.11, colw=1.95,
                      marker=None, top=1.98, list_y=3.42, step=0.385):
    """Course timeline with every session listed under its block.

    cols: dicts with kicker, title, date, meta, optional tag, and
          sessions = [(number, name, date, is_ours), ...]
    """
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(axis_y), Inches(13.333), Pt(1))
    _fill(line, "3A3D52")
    _noline(line)
    for i, c in enumerate(cols):
        x = x0 + i * pitch
        tick = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.02), Inches(axis_y - 0.10),
                                  Pt(1), Inches(0.22))
        _fill(tick, "3A3D52")
        _noline(tick)
        txt(s, c["kicker"], x, top, colw + 0.1, 0.22, HNL, 10.5, GREY_TXT)
        txt(s, c["title"], x, top + 0.22, colw, 0.46, HN, 13, NAVY, bold=True, spacing=1.02)
        txt(s, c["date"], x, axis_y + 0.14, colw + 0.25, 0.26, HNL, 12.5, PURPLE)
        txt(s, c["meta"], x, axis_y + 0.40, colw + 0.25, 0.20, HNL, 9, GREY_TXT)
        for k, (num, name, when, ours) in enumerate(c["sessions"]):
            y = list_y + k * step
            dot = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                     Inches(0.30), Inches(0.21))
            _round(dot, 20000)
            _fill(dot, PURPLE if ours else GREY1)
            _noline(dot)
            txt(s, str(num), x, y, 0.30, 0.21, HN, 8.5, WHITE if ours else "8A93A0",
                bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            box = s.shapes.add_textbox(Inches(x + 0.36), Inches(y - 0.045),
                                       Inches(colw - 0.34), Inches(0.36))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            par = tf.paragraphs[0]
            par.line_spacing = 1.05
            _run(par, name, HN if ours else HNL, 9, NAVY, bold=ours)
            if when:
                _run(par, "  " + when, HNL, 9, GREY_TXT)
    if marker is not None:
        mx = x0 + marker * pitch
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(mx - 0.10), Inches(axis_y - 0.13),
                                 Inches(0.26), Inches(0.26))
        _fill(dot, PURPLE)
        _line(dot, WHITE, 1.75)
